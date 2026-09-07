from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook
import os

class Scraper:
    def __init__(self, url, filename):
        self.url = url
        self.filename = filename
        self.wd = None
        self.ws = None
        self.carUrl = None
        self.model = None
        self.generation = None
        self.cost = None
        self.mileage = None
        self.fuelType = None
        self.gearbox = None
        self.location = None
        self.shortDescription = None
        self.time = None

    def setupXL(self):
        if os.path.exists(self.filename):
            self.wd = load_workbook(self.filename)
            self.ws = self.wd.active
        else:
            self.wd = Workbook()
            self.ws = self.wd.active
            self.ws.append(["URL", "Model", "Generation", "Cost $", "Mileage", "Fuel Type", "Gearbox", "Location", "Time"])

    def addCarToXL(self):
        current_row = self.ws.max_row + 1
        carData = [self.carUrl, self.model, self.generation, self.cost, self.mileage, self.fuelType, self.gearbox, self.location, self.time]
        for col,value in enumerate(carData, start=1):
            self.ws.cell(row=current_row, column=col).value = value

    def scrape(self):
        self.setupXL()
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            page = browser.new_page()
            page.goto(self.url)
            page.wait_for_timeout(1000)

            try:
                while True:
                    cars = page.query_selector_all(".link.product-card.horizontal")
                    for car in cars:
                        self.carUrl = "https://auto.ria.com" + car.get_attribute("href")
                        self.model = car.query_selector(".common-text.size-16-20.titleS.fw-bold.mb-4").inner_text()
                        generationEl = car.query_selector(".common-text.size-14-16.ellipsis-1.mb-8")
                        self.generation = generationEl.inner_text() if generationEl else None
                        self.cost = car.query_selector(".common-text.titleM.c-green").inner_text().replace('\xa0', ' ').strip()
                        elements = car.query_selector_all(".common-text.ellipsis-1.body")
                        self.mileage = elements[0].inner_text()
                        self.fuelType = elements[1].inner_text()
                        self.gearbox = elements[2].inner_text()
                        self.location = elements[3].inner_text()
                        shortDescriptionEl = car.query_selector(".common-text.footnote.mt-12")
                        self.shortDescription = shortDescriptionEl.inner_text() if shortDescriptionEl else None
                        timeEl = car.query_selector(".common-text.footnote.c-contrastSecondary")
                        self.time = timeEl.inner_text() if timeEl else None
                        self.addCarToXL()
                        print(f"Saved: {self.carUrl, self.model, self.generation, self.cost, self.mileage, self.fuelType, self.gearbox, self.location, self.shortDescription, self.time}")
                    button = page.query_selector(".ml-16 > .size-large")
                    if button:
                        button.click()
                        page.wait_for_timeout(5000)
                    else:
                        break
            except Exception as e:
                print(f'An error occurred during scraping: {e}')

            finally:
                print(f'Saving data to: {self.filename}')
                self.wd.save(self.filename)
                browser.close()

autoRia = Scraper("https://auto.ria.com/uk/search/?search_type=1&category=1&all[0].any[0].state=16&all[0].any[0].any[0].city=16&abroad=0&customs_cleared=1", "autoRia.xlsx")
autoRia.scrape()